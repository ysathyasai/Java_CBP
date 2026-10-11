# -*- coding: utf-8 -*-
"""
Report helper formatting functions and XML styling for YieldGrid Academic Project Report.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def add_code_block(doc, code_str, title=None):
    if title:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(1.5)
        p_title.paragraph_format.keep_with_next = True
        r_title = p_title.add_run(f"Source Excerpt: {title}")
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(16, 42, 67)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F8F9FA")
    set_cell_margins(cell, top=45, bottom=45, left=120, right=120)
    set_cell_border(cell,
                    left={'val': 'single', 'sz': 16, 'color': '1E3A8A'},
                    top={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                    right={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                    bottom={'val': 'single', 'sz': 4, 'color': 'CBD5E1'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.02
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.2)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3.5)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.5)
    run.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(15.5)
    run.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_body_p(doc, text, bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_figure(doc, img_path, caption, width=Inches(5.5)):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(img_path, width=width)
        add_caption(doc, caption)
    else:
        print(f"Warning: Image file not found: {img_path}")

def add_side_by_side_figures(doc, img_path1, cap1, img_path2, cap2, width=Inches(3.1)):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    pair = [(img_path1, cap1), (img_path2, cap2)]
    for col_idx, (img, cap) in enumerate(pair):
        cell_img = tbl.cell(0, col_idx)
        cell_img.width = Inches(3.25)
        set_cell_margins(cell_img, top=10, bottom=5, left=15, right=15)
        p_img = cell_img.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        if os.path.exists(img):
            r = p_img.add_run()
            r.add_picture(img, width=width)

        cell_cap = tbl.cell(1, col_idx)
        cell_cap.width = Inches(3.25)
        set_cell_margins(cell_cap, top=0, bottom=20, left=15, right=15)
        p_cap = cell_cap.paragraphs[0]
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(1)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(cap)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(8.5)
        r_cap.bold = True
        r_cap.font.color.rgb = RGBColor(30, 41, 59)

    for r in tbl.rows:
        for c in r.cells:
            tcPr = c._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(tcBorders)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

print("Report helper modules initialized.")
